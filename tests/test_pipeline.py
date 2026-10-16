"""Проверки реального HTTP-сбора, сохранности данных и учёта ошибок."""
import json
import pytest
from openpyxl import load_workbook
from web_pipeline.backends import FetchResult
from web_pipeline.config import load_config
from web_pipeline.pipeline import run
from web_pipeline.storage import write_xlsx, CELL_LIMIT


def config(site, engines=None):
    return {"engines": engines or ["curl_cffi"], "sources": [
        {"id": "local", "url": site, "title_selector": "h1", "text_selector": "main"}],
        "timeout_seconds": 15, "delay_seconds": 0, "retry_delay_seconds": 0,
        "retries": 1, "headless": True}


def test_http_end_to_end(site, tmp_path):
    result = run(config(site), tmp_path / "run")
    assert result["successful_pairs"] == 1
    assert result["failed_pairs"] == 0
    book = load_workbook(tmp_path / "run/data/processed.xlsx")
    rows = list(book.active.values)
    record = dict(zip(rows[0], rows[1]))
    assert record["title"] == "Научные данные"
    assert record["text"] == "Проверяемый текст для сбора."
    assert record["word_count"] == 4
    manifest = json.loads((tmp_path / "run/logs/manifest.json").read_text(encoding="utf-8"))
    assert manifest["artifacts"]["data/raw.xlsx"]


def test_blocked_is_recorded_without_retry(site, tmp_path):
    cfg = config(site + "/blocked")
    result = run(cfg, tmp_path / "run")
    assert result["failed_pairs"] == 1
    attempts = json.loads((tmp_path / "run/logs/attempts.json").read_text())
    assert len(attempts) == 1
    assert attempts[0]["status"] == 403


def test_retry_and_deduplication(site, tmp_path):
    calls = []
    def fetcher(engine, source, cfg):
        calls.append(engine)
        if len(calls) == 1:
            return FetchResult(site, 503, "temporary")
        return FetchResult(site, 200, "<h1>Название</h1><main>Общий текст</main>")
    result = run(config(site, ["curl_cffi", "patchright"]), tmp_path / "run", fetcher)
    assert len(calls) == 3
    assert result["successful_pairs"] == 2
    assert result["duplicate_records"] == 1
    assert result["unique_nonempty_records"] == 1


def test_empty_selection_and_missing_backend(site, tmp_path):
    def fetcher(engine, source, cfg):
        if engine == "camoufox":
            raise ModuleNotFoundError("private information must not be logged")
        return FetchResult(site, 200, "<p>Без нужного селектора</p>")
    result = run(config(site, ["curl_cffi", "camoufox"]), tmp_path / "run", fetcher)
    assert result["failed_pairs"] == 2
    assert result["empty_text"] == 1
    assert "private information" not in (tmp_path / "run/logs/attempts.json").read_text()


def test_excel_formula_and_long_text_roundtrip(tmp_path):
    long_text = "А" * (CELL_LIMIT * 2 + 7)
    path = tmp_path / "data.xlsx"
    write_xlsx(path, {"raw": [{"title": "=1+1", "html": long_text, "text": "a\x00b"}]})
    book = load_workbook(path)
    assert book["raw"]["A2"].data_type == "s"
    assert book["raw"]["A2"].value == "=1+1"
    assert book["raw"]["C2"].value == r"a\u0000b"
    assert "".join(row[4] for row in list(book["long_values"].values)[1:]) == long_text


@pytest.mark.parametrize("edit", [
    {"engines": ["unknown"]}, {"timeout_seconds": -1}, {"delay_seconds": float("inf")},
    {"retries": True}, {"sources": []},
    {"sources": [{"id": "x", "url": "file:///tmp/test"}]},
    {"sources": [{"id": "x", "url": "https://example.com", "text_selector": "["}]},
])
def test_invalid_config_rejected_before_network(tmp_path, site, edit):
    cfg = config(site)
    cfg.update(edit)
    path = tmp_path / "config.json"
    path.write_text(json.dumps(cfg), encoding="utf-8")
    with pytest.raises(ValueError):
        load_config(path)


def test_existing_run_is_preserved(site, tmp_path):
    with pytest.raises(FileExistsError):
        run(config(site), tmp_path)
