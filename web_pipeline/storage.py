"""XLSX без преобразования внешнего текста в формулы и без потери длинных полей."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

CELL_LIMIT = 30000


def write_xlsx(path: Path, sheets: dict[str, list[dict]]) -> None:
    book = Workbook()
    book.remove(book.active)
    overflow = []
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        keys = list(dict.fromkeys(key for row in rows for key in row)) or ["сообщение"]
        sheet.append(keys)
        for row_index, row in enumerate(rows, 2):
            for col_index, key in enumerate(keys, 1):
                value = row.get(key, "")
                if isinstance(value, str):
                    # Недопустимые XML-символы заменяются печатными escape-последовательностями.
                    value = ILLEGAL_CHARACTERS_RE.sub(lambda m: f"\\u{ord(m[0]):04x}", value)
                    if len(value) > CELL_LIMIT:
                        for part, start in enumerate(range(0, len(value), CELL_LIMIT), 1):
                            overflow.append({"sheet": name, "row": row_index, "field": key,
                                             "part": part, "value": value[start:start + CELL_LIMIT]})
                        value = "Содержимое сохранено частями на листе long_values"
                cell = sheet.cell(row_index, col_index, value)
                if isinstance(value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="Times New Roman", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24445C")
        for column in sheet.columns:
            key = column[0].value
            sheet.column_dimensions[column[0].column_letter].width = 70 if key in {"text", "html"} else 24
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Times New Roman", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    if overflow:
        sheet = book.create_sheet("long_values")
        keys = list(overflow[0])
        sheet.append(keys)
        for record in overflow:
            sheet.append([record[key] for key in keys])
            sheet.cell(sheet.max_row, 5).data_type = "s"
        sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.xlsx")
    book.save(temporary)
    temporary.replace(path)
